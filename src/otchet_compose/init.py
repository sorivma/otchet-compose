"""Interactive ``otchet-compose init`` command.

Asks a series of questions and writes a starter ``otchet-compose.yml``
that is ready to use with ``otchet-compose gen``.
"""

from datetime import date
import os
from pathlib import Path
import tempfile

import yaml
from .generator.title_page import list_available_templates, _collect_placeholders, BUNDLED_TEMPLATES_DIR, _USER_TEMPLATES_DIR
from docx import Document as _DocxDocument


# Human-readable hints for well-known template params
_PARAM_HINTS: dict[str, str] = {
    "discipline":  "Название дисциплины",
    "work_number": "Номер работы, например: № 4",
    "work_title":  "Название работы",
    "work_type":   "Тип работы, например: ОТЧЁТ ПО ЛАБОРАТОРНОЙ РАБОТЕ",
    "student":     "ФИО студента, например: Иванов И.И.",
    "group":       "Номер группы, например: ИКБО-01-22",
    "teacher":     "ФИО преподавателя, например: Петров П.П.",
    "year":        "Год",
    "institute":   "Название института",
    "department":  "Название кафедры",
}

_STARTER_CONTENT = """\
content:
  - type: heading
    text: "Цель работы"
    level: 1
    structural: true

  - type: paragraph
    text: >
      Опишите цель лабораторной работы.

  - type: heading
    text: "Ход работы"
    level: 1
    structural: true

  - type: heading
    text: "Постановка задачи"
    level: 2
    structural: false

  - type: paragraph
    text: >
      Опишите постановку задачи.

  - type: heading
    text: "Заключение"
    level: 1
    structural: true

  - type: paragraph
    text: >
      Подведите итоги работы.
"""

_STARTER_BLOCKS = yaml.safe_load(_STARTER_CONTENT)["content"]


def init_command(args) -> int:
    """Execute the ``init`` subcommand."""
    print("Инициализация конфигурации otchet-compose\n")

    # --- output config path ---
    config_path = Path(_ask("Путь к файлу конфигурации", "otchet-compose.yml"))
    if config_path.exists():
        overwrite = _ask_bool(f"Файл {config_path} уже существует. Перезаписать?", default=False)
        if not overwrite:
            print("Отменено.")
            return 0

    # --- document output ---
    output = _ask("Путь к выходному DOCX файлу", "./build/report.docx")
    toc = _ask_bool("Включить оглавление?", default=True)

    # --- title page ---
    use_title_page = _ask_bool("Добавить титульный лист из шаблона?", default=True)

    title_page_yaml = ""
    reserve_yaml = "  reserve_title_page: false"

    if use_title_page:
        title_page_yaml = _collect_title_page_yaml()
        reserve_yaml = ""
    else:
        reserve = _ask_bool("Зарезервировать пустую первую страницу под титульник?", default=False)
        if reserve:
            reserve_yaml = "  reserve_title_page: true"

    # --- assemble YAML ---
    lines = ["version: 1", "", "document:"]
    lines.append(f'  output: "{output}"')
    lines.append(f"  toc: {'true' if toc else 'false'}")
    if reserve_yaml:
        lines.append(reserve_yaml)
    if title_page_yaml:
        lines.append(title_page_yaml)
    lines.append("")
    lines.append(_STARTER_CONTENT)

    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"\nКонфигурация сохранена: {config_path}")
    print(f"Запустите генерацию: otchet-compose gen -f {config_path}")
    return 0


def create_config(
    config_path: str | Path,
    *,
    document_output: str = "./build/report.docx",
    toc: bool = True,
    template: str | None = None,
    params: dict[str, str] | None = None,
    reserve_title_page: bool = False,
    force: bool = False,
) -> Path:
    """Create a starter config non-interactively and atomically."""
    path = Path(config_path).resolve()
    if path.exists() and not force:
        raise FileExistsError(f"Config already exists; use --force to overwrite: {path}")
    if template and template not in list_available_templates():
        available = ", ".join(list_available_templates())
        raise ValueError(f"Unknown title page template: {template}. Available templates: {available}")

    document: dict[str, object] = {"output": document_output, "toc": toc}
    if template:
        document["title_page"] = {"template": template, "params": params or {}}
    else:
        document["reserve_title_page"] = reserve_title_page
    payload = {"version": 1, "document": document, "content": _STARTER_BLOCKS}

    path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent, delete=False)
    handle.close()
    temp = Path(handle.name)
    try:
        temp.write_text(yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8")
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)
    return path


# ---------------------------------------------------------------------------
# Title page section builder
# ---------------------------------------------------------------------------

def _collect_title_page_yaml() -> str:
    """Interactively ask for a template and its params; return YAML snippet."""
    available = list_available_templates()
    if not available:
        print("Встроенные шаблоны не найдены. Раздел title_page будет пропущен.")
        return ""

    print(f"\nДоступные шаблоны: {', '.join(available)}")
    default_tmpl = available[0]
    template = _ask("Шаблон", default_tmpl)

    if template not in available:
        print(f"Шаблон '{template}' не найден, будет использован '{default_tmpl}'.")
        template = default_tmpl

    params = _ask_template_params(template)

    lines = ["  title_page:"]
    lines.append(f"    template: {template}")
    if params:
        lines.append("    params:")
        for key, value in params.items():
            lines.append(f'      {key}: "{value}"')
    return "\n".join(lines)


def _ask_template_params(template: str) -> dict[str, str]:
    """Load the template .docx, collect its placeholders, and ask for each value."""
    tmpl_path = (_USER_TEMPLATES_DIR / f"{template}.docx")
    if not tmpl_path.exists():
        tmpl_path = BUNDLED_TEMPLATES_DIR / f"{template}.docx"
    if not tmpl_path.exists():
        return {}

    tmpl_doc = _DocxDocument(str(tmpl_path))
    placeholders = sorted(_collect_placeholders(tmpl_doc))

    if not placeholders:
        return {}

    print(f"\nПараметры шаблона «{template}» (Enter — оставить пустым):")
    params: dict[str, str] = {}
    for key in placeholders:
        hint = _PARAM_HINTS.get(key, key)
        default = str(date.today().year) if key == "year" else ""
        value = _ask(f"  {hint} [{key}]", default)
        params[key] = value

    return params


# ---------------------------------------------------------------------------
# Prompt helpers
# ---------------------------------------------------------------------------

def _ask(prompt: str, default: str = "") -> str:
    """Print *prompt* with an optional default and return the user's input."""
    if default:
        display = f"{prompt} [{default}]: "
    else:
        display = f"{prompt}: "
    answer = input(display).strip()
    return answer if answer else default


def _ask_bool(prompt: str, default: bool = True) -> bool:
    """Ask a yes/no question and return a bool."""
    hint = "[Y/n]" if default else "[y/N]"
    answer = input(f"{prompt} {hint}: ").strip().lower()
    if not answer:
        return default
    return answer in ("y", "yes", "да", "д")
